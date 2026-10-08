from pathlib import Path
from io import BytesIO
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "output"
ROWS = [('来源文件', '页面内容', '转换后来源页码', '合并PDF页码', '归属及保留说明'), ('4-2026年第三次临时股东会通知回执.doc', '通知回执', 7, 1, '确认人王强（原文含空格）'), ('7-2026年第三次临时股东会签到表（拆分）.doc', '签到页', 1, 2, '股东名称王强；整页含其他股东'), ('9-2026年第三次临时股东会表决票.doc', '表决票正文', 1, 3, '股东名称王强，持股数528.8246万股'), ('9-2026年第三次临时股东会表决票.doc', '表决票签署页', 2, 4, '紧接王强表决票正文；签署页未单独标注姓名'), ('11-2026年第三次临时股东会会议决议.doc', '股东签署页', 12, 5, '含王强个人及其代表合伙企业的签署位置'), ('11-2026年第三次临时股东会会议决议.doc', '股东签署页续页', 13, 6, '承接前页合伙企业签字及日期栏'), ('12-2026年第三次临时股东会会议记录.doc', '股东签署页', 14, 7, '含王强个人及其代表合伙企业的签署位置')]
REPORT = '# 王强股东会签字页提取说明\n\n## 整理结果\n\n根据用户上传的七份文件，按文件编号和原页序提取王强作为股东的签字相关页面，合并为7页PDF。原始文件未修改。\n\n页码指本次Word文件转换为PDF后的页码，并非保证与本机WPS或Word显示的页码一致。\n\n## 提取清单\n\n|来源文件编号|页面内容|转换后来源页码|合并PDF页码|\n|---|---|---|---|\n|4|通知回执|7|1|\n|7|签到页|1|2|\n|9|表决票正文|1|3|\n|9|表决票签署页|2|4|\n|11|股东签署页|12|5|\n|11|股东签署页续页|13|6|\n|12|股东签署页|14|7|\n\n## 范围及核对说明\n\n- 表决票签署页没有单独标注王强姓名，依据紧邻的王强表决票正文确认归属，并保留正文供核对。\n- 签到页、决议及记录股东签署页包含其他股东或合伙企业的信息，保留整页，未裁切。决议签署页因转换分页产生续页，已一并保留。\n- 会议通知没有王强股东签署位置；授权委托书为未填写委托人名称的通用模板，未纳入。\n- 王强以董事、高管或会议主持人身份出现的其他签署页，未纳入此次以股东身份为范围的汇总。\n- 文件显示的是待签署位置，本次整理不证明已经实际签字，也不构成对会议程序或文件内容的法律合规审查。\n\n## 输出文件\n\n- `report.md`：提取说明。\n- `data.xlsx`：来源文件、来源页码及合并页码清单。\n- `王强-2026年第三次临时股东会签字页汇总.pdf`：合并PDF。\n'

def main():
    OUTPUT.mkdir(exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "签字页清单"
    for row in ROWS:
        sheet.append(row)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="24527A")
    for column, width in [("A", 60), ("B", 24), ("C", 22), ("D", 18), ("E", 65)]:
        sheet.column_dimensions[column].width = width
    for row in sheet.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for index in range(2, len(ROWS) + 1):
        sheet.row_dimensions[index].height = 36
    buffer = BytesIO()
    workbook.save(buffer)
    assert list(load_workbook(BytesIO(buffer.getvalue())).active.values) == ROWS
    (OUTPUT / "data.xlsx").write_bytes(buffer.getvalue())
    (OUTPUT / "report.md").write_text(REPORT, encoding="utf-8")
    print("已生成 output/report.md 和 output/data.xlsx（7条页面记录）")

if __name__ == "__main__":
    main()
